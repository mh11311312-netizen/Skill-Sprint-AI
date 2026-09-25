"""Step 6 - Text extraction with Python.
Every extracted line keeps its location: page number (PDF) or paragraph number (DOCX/TXT),
plus a 'hidden' flag for text a human cannot see (white or tiny font) - a common prompt-injection trick."""
import io
import re
import pdfplumber
import docx
from docx.shared import Pt


def _pdf_color_is_white(color):
    if color is None:
        return False
    vals = color if isinstance(color, (list, tuple)) else [color]
    try:
        vals = [float(v) for v in vals]
    except (TypeError, ValueError):
        return False
    if len(vals) == 4:                       # CMYK white = 0,0,0,0
        return all(v == 0 for v in vals)
    return all(v >= 0.95 for v in vals)      # gray 1 or RGB 1,1,1


def parse_pdf(data):
    lines = []
    with pdfplumber.open(io.BytesIO(data)) as pdf:
        for page_no, page in enumerate(pdf.pages, start=1):
            for ln in page.extract_text_lines(return_chars=True):
                chars = [c for c in ln["chars"] if c["text"].strip()]
                hidden_chars = sum(1 for c in chars if c["size"] < 3 or _pdf_color_is_white(c.get("non_stroking_color")))
                lines.append({"text": ln["text"].strip(), "page": page_no, "para": None, "style": "",
                              "hidden": bool(chars) and hidden_chars / len(chars) > 0.5})
    return {"lines": [l for l in lines if l["text"]], "tables": []}


def _run_hidden(run):
    f = run.font
    if f.hidden:
        return True
    if f.size is not None and f.size < Pt(3):
        return True
    try:
        if f.color is not None and f.color.rgb is not None and str(f.color.rgb).upper() in ("FFFFFF", "FEFEFE"):
            return True
    except (AttributeError, ValueError):
        pass
    return False


def parse_docx(data):
    d = docx.Document(io.BytesIO(data))
    lines = []
    for i, p in enumerate(d.paragraphs, start=1):
        text = p.text.strip()
        if not text:
            continue
        runs = [r for r in p.runs if r.text.strip()]
        hidden = bool(runs) and all(_run_hidden(r) for r in runs)
        style = p.style.name if p.style is not None else ""
        lines.append({"text": text, "page": None, "para": i, "style": style, "hidden": hidden})
    tables = [[[c.text.strip() for c in row.cells] for row in t.rows] for t in d.tables]
    return {"lines": lines, "tables": tables}


def parse_text(data):
    text = data.decode("utf-8-sig", errors="replace")
    lines = []
    for i, raw in enumerate(text.splitlines(), start=1):
        raw = raw.strip()
        if not raw:
            continue
        hidden = "color:#ffffff" in raw.replace(" ", "").lower()
        raw = re.sub(r"<[^>]+>", "", raw).strip()          # strip html tags from markdown
        raw = re.sub(r"^#+\s*", "", raw)                  # markdown headings
        lines.append({"text": raw, "page": None, "para": i, "style": "", "hidden": hidden})
    return {"lines": lines, "tables": []}


def parse_file(data, ext):
    if ext == "pdf":
        return parse_pdf(data)
    if ext == "docx":
        return parse_docx(data)
    return parse_text(data)
